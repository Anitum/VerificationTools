#!/usr/bin/env python3
"""
annotation_excel_gen.py

Generates per-channel annotation Excel files from ROI overlay images.
Each row = one representative frame per 5-second bucket.

Frame timestamp = clip_start + (frame_id - 1) × (VID_STRIDE / FPS)
                = clip_start + (frame_id - 1) × 1.5s   (stride=6, fps=4)

5s bucket = floor(frame_timestamp / 5s) × 5s

One Excel file is written per channel. Annotators fill in the verified
class counts (F–L) by looking at the bounding-box overlays drawn on each image.

Usage:
    python3 annotation_excel_gen.py
    python3 annotation_excel_gen.py --input-dir <path> --output-dir <path>
"""

import re
import math
import io
import argparse
from pathlib import Path
from datetime import datetime, timedelta
from collections import defaultdict

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage


# ── Configuration ──────────────────────────────────────────────────────────────

VID_STRIDE   = 6
FPS          = 4
FRAME_STEP_S = VID_STRIDE / FPS   # 1.5 s between consecutive processed frames
BUCKET_S     = 5                   # headcount CSV is binned at 5-second intervals

THUMB_W      = 240                 # thumbnail width  (px)
THUMB_H      = 135                 # thumbnail height (px) – 16:9
ROW_H_PTS    = 104                 # Excel row height in points (~140 px on screen)

DEFAULT_INPUT  = Path("/home/aabadmin/SHUBHAM/video_detector/verification_tool/roi_overlay_output")
DEFAULT_OUTPUT = Path("/home/aabadmin/SHUBHAM/video_detector/verification_tool/annotation_excel")

# Detection class labels (class_0 … class_6)
CLASS_LABELS = ["customer", "manager", "waiter",
                "sweets_people", "back_kitchen", "cleaner", "phone"]

# ── Regex patterns ─────────────────────────────────────────────────────────────

FOLDER_RE = re.compile(r"^(.+)_(channel\d+)_(\d{8}T\d{6})_frames$")
IMAGE_RE  = re.compile(r"^.+?_(\d+)\.jpg$", re.IGNORECASE)


# ── Parsing helpers ────────────────────────────────────────────────────────────

def parse_folder(name: str):
    """Return metadata dict from a *_frames directory name, or None."""
    m = FOLDER_RE.match(name)
    if not m:
        return None
    outlet, channel, ts = m.groups()
    return {
        "outlet":     outlet,
        "channel":    channel,
        "clip_start": datetime.strptime(ts, "%Y%m%dT%H%M%S"),
    }


def frame_id_from_name(name: str):
    """Extract the integer frame-id from an image filename."""
    m = IMAGE_RE.match(name)
    return int(m.group(1)) if m else None


def to_bucket(clip_start: datetime, frame_id: int):
    """
    Compute:
      frame_ts = clip_start + (frame_id - 1) × FRAME_STEP_S
      bucket   = floor(frame_ts to nearest BUCKET_S seconds)
    Returns (bucket_datetime, frame_datetime).
    """
    frame_ts  = clip_start + timedelta(seconds=(frame_id - 1) * FRAME_STEP_S)
    epoch     = math.floor(frame_ts.timestamp() / BUCKET_S) * BUCKET_S
    bucket_dt = datetime.fromtimestamp(epoch)
    return bucket_dt, frame_ts


def representative_frames(folder: Path, meta: dict) -> dict:
    """
    For each 5-second bucket that this clip covers, return the single image
    whose frame_ts is closest to the bucket start time.

    Returns: { bucket_datetime: (img_path, frame_datetime) }
    """
    best = {}  # bucket_dt → (img_path, frame_ts, distance_seconds)

    for img in sorted(folder.glob("*.jpg")):
        fid = frame_id_from_name(img.name)
        if fid is None:
            continue
        bucket_dt, frame_ts = to_bucket(meta["clip_start"], fid)
        dist = abs((frame_ts - bucket_dt).total_seconds())
        if bucket_dt not in best or dist < best[bucket_dt][2]:
            best[bucket_dt] = (img, frame_ts, dist)

    return {b: (p, ts) for b, (p, ts, _) in best.items()}


# ── Thumbnail helper ───────────────────────────────────────────────────────────

def make_thumb(img_path: Path) -> io.BytesIO:
    """Resize image to THUMB_W × THUMB_H and return as in-memory JPEG buffer."""
    img = PILImage.open(img_path).convert("RGB")
    img = img.resize((THUMB_W, THUMB_H), PILImage.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=82)
    buf.seek(0)
    return buf


# ── Excel styles ───────────────────────────────────────────────────────────────

HDR_FILL    = PatternFill("solid", fgColor="1F4E79")
HDR_FONT    = Font(color="FFFFFF", bold=True, size=10)
ALT_FILL    = PatternFill("solid", fgColor="D6E4F0")
CTR         = Alignment(horizontal="center", vertical="center")
WRAP_CTR    = Alignment(horizontal="center", vertical="center", wrap_text=True)
VERIFY_FILL = PatternFill("solid", fgColor="FFF2CC")   # pale yellow — annotator input

thin   = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# Column definitions: (header label, column width)
# A=Frame  B=Channel  C=Computed Timestamp  D=Frame Time  E=Clip Name
# F-L = annotator counts   M = Notes
COLUMNS = (
    [("Frame",                       34)],
    [("Channel",                     14)],
    [("Computed Timestamp (T2)",     24)],
    [("Frame Time",                  22)],
    [("Clip Name",                   44)],
    [(f"✓ {lbl}", 13) for lbl in CLASS_LABELS],   # F–L  annotator input
    [("Notes",                       32)],         # M
)
HEADERS = [item for group in COLUMNS for item in group]

N_CLASS          = len(CLASS_LABELS)
VERIFY_START_COL = 6                               # F
VERIFY_END_COL   = VERIFY_START_COL + N_CLASS - 1 # L
NOTES_COL        = VERIFY_END_COL + 1             # M


# ── Excel writer ───────────────────────────────────────────────────────────────

def write_channel_excel(channel: str, clips: list, output_path: Path) -> int:
    """
    Write one Excel workbook for `channel`.
    clips: list of (folder_path, meta_dict), sorted by clip_start.
    Returns number of data rows written.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = channel[:31]

    # ── Header row ──
    for ci, (label, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=ci, value=label)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CTR
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "B2"

    # Annotator column headers: amber
    for ci in range(VERIFY_START_COL, NOTES_COL + 1):
        cell = ws.cell(row=1, column=ci)
        cell.fill = PatternFill("solid", fgColor="F4B942")
        cell.font = Font(color="1F1F1F", bold=True, size=10)

    # ── Data rows ──
    row = 2
    for folder, meta in sorted(clips, key=lambda x: x[1]["clip_start"]):
        frames = representative_frames(folder, meta)

        for bucket_dt in sorted(frames):
            img_path, frame_ts = frames[bucket_dt]
            use_alt = (row % 2 == 0)

            # Embed thumbnail in column A
            xl_img        = XLImage(make_thumb(img_path))
            xl_img.width  = THUMB_W
            xl_img.height = THUMB_H
            ws.add_image(xl_img, f"A{row}")

            row_values = [
                None,                                        # A – image
                meta["channel"],                             # B
                bucket_dt.strftime("%Y-%m-%d %H:%M:%S"),    # C
                frame_ts.strftime("%Y-%m-%d %H:%M:%S"),      # D
                folder.name,                                 # E
                *[""] * N_CLASS,                             # F–L  annotator counts
                "",                                          # M   notes
            ]

            for ci, val in enumerate(row_values, start=1):
                if ci == 1:
                    continue  # image cell
                cell = ws.cell(row=row, column=ci, value=val if val != "" else None)
                cell.alignment = WRAP_CTR if ci in (3, 5) else CTR
                cell.border = BORDER

                # Alternating row background for non-input columns
                if use_alt and not (VERIFY_START_COL <= ci <= NOTES_COL):
                    cell.fill = ALT_FILL

                # Annotator input columns: always pale yellow
                if VERIFY_START_COL <= ci <= NOTES_COL:
                    cell.fill = VERIFY_FILL

            ws.row_dimensions[row].height = ROW_H_PTS
            row += 1

    # ── Legend sheet ──
    ls = wb.create_sheet("Legend")
    ls["A1"] = "Column"
    ls["B1"] = "Description"
    for ci in (1, 2):
        ls.cell(1, ci).font = Font(bold=True)

    legend_rows = [
        ("Frame",
         "Thumbnail of the representative frame for this 5-second window. "
         "The CCTV-burned timestamp visible in the image is T1 (camera time, IST)."),
        ("Channel",
         "Physical camera channel ID"),
        ("Computed Timestamp (T2)",
         "T2 = clip_start (UTC, from filename) + (frame_id - 1) × 1.5 s, floored to 5 s. "
         "NOTE: T2 is UTC. The camera OSD (T1) shows IST (UTC+5:30), so T1 will appear "
         "~5 h 30 min ahead of T2. This is expected."),
        ("Frame Time",
         "Exact computed timestamp of this frame (before 5-second bucketing)"),
        ("Clip Name",
         "Source video clip folder name"),
        *[(f"✓ {lbl}",
           f"ANNOTATOR: enter the count of '{lbl}' visible in the frame")
          for lbl in CLASS_LABELS],
        ("Notes",
         "Free-text observations"),
    ]
    for r, (col, desc) in enumerate(legend_rows, start=2):
        ls.cell(r, 1).value = col
        ls.cell(r, 2).value = desc
    ls.column_dimensions["A"].width = 26
    ls.column_dimensions["B"].width = 80

    wb.save(output_path)
    return row - 2


# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Generate per-channel annotation Excel files from ROI overlay images"
    )
    parser.add_argument("--input-dir",  default=str(DEFAULT_INPUT),
                        help="Directory containing *_frames subdirectories")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT),
                        help="Output directory for Excel files")
    args = parser.parse_args()

    input_dir  = Path(args.input_dir)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    if not input_dir.is_dir():
        print(f"ERROR: input directory not found: {input_dir}")
        return

    # Group clips by channel
    by_channel = defaultdict(list)
    skipped = 0
    for folder in sorted(input_dir.iterdir()):
        if not folder.is_dir():
            continue
        meta = parse_folder(folder.name)
        if meta:
            by_channel[meta["channel"]].append((folder, meta))
        else:
            skipped += 1

    total_clips = sum(len(v) for v in by_channel.values())
    print(f"Input : {input_dir}")
    print(f"Output: {output_dir}")
    print(f"Found : {total_clips} clips across {len(by_channel)} channels"
          + (f" ({skipped} dirs skipped)" if skipped else ""))
    print()

    total_rows = 0
    for channel, clips in sorted(by_channel.items()):
        out_path = output_dir / f"annotation_{channel}.xlsx"
        print(f"  {channel:20s}  {len(clips):3d} clips  →  {out_path.name} ...", end=" ", flush=True)
        n = write_channel_excel(channel, clips, out_path)
        print(f"{n} rows")
        total_rows += n

    print(f"\nDone.  {total_rows} total rows written to: {output_dir}")


if __name__ == "__main__":
    main()
How 